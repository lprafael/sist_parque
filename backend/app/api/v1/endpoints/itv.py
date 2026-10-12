import io
from datetime import date
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, and_, update, or_
from sqlalchemy.orm import selectinload

from app.core.database import get_db
from app.core.security import get_current_user, require_roles
from app.models import ItvBus, Bus, Eot
from app.services.sistema_logs import registrar_auditoria
from app.schemas import ItvBusCreate, ItvBusUpdate, ItvBusOut, HistorialItvOut

router = APIRouter(prefix="/itv", tags=["ITV - Inspección Técnica"])


def estado_itv(venc: Optional[date]) -> tuple[str, Optional[int]]:
    if not venc:
        return "VENCIDO", None
    diff = (venc - date.today()).days
    if diff < 0:
        estado = "VENCIDO"
    elif diff <= 7:
        estado = "CRITICO"
    elif diff <= 30:
        estado = "POR_VENCER"
    else:
        estado = "VIGENTE"
    return estado, diff


@router.get("", response_model=dict)
async def listar_itv(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    id_bus: Optional[int] = None,
    estado: Optional[str] = None,
    vence_antes_de: Optional[date] = None,
    solo_vigentes: bool = Query(True, description="Si es True, sólo retorna los registros de ITV vigentes"),
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user)
):
    q = select(ItvBus)
    filters = []
    if solo_vigentes:
        filters.append(ItvBus.es_vigente == True)
    if id_bus:
        filters.append(ItvBus.id_bus == id_bus)
    if vence_antes_de:
        filters.append(ItvBus.fecha_vencimiento <= vence_antes_de)
    if filters:
        q = q.where(and_(*filters))

    total = (await db.execute(select(func.count()).select_from(q.subquery()))).scalar()
    itv_list = (
        await db.execute(q.offset((page - 1) * page_size).limit(page_size).order_by(ItvBus.fecha_vencimiento))
    ).scalars().all()

    items = []
    for itv in itv_list:
        est, dias = estado_itv(itv.fecha_vencimiento)
        if estado and est != estado.upper():
            continue
        items.append(ItvBusOut(
            **{c.name: getattr(itv, c.name) for c in ItvBus.__table__.columns},
            dias_para_vencer=dias,
            estado_itv=est
        ))

    return {"total": total, "page": page, "page_size": page_size, "items": items}


@router.get("/excel")
async def exportar_itv_excel(
    estado: Optional[str] = Query(None, description="Filtrar por estado ITV: VIGENTE, POR_VENCER, CRITICO, VENCIDO"),
    search: Optional[str] = Query(None, description="Buscar por RUA o N° certificado"),
    solo_vigentes: bool = Query(True, description="Si es True, sólo exporta registros ITV vigentes"),
    db: AsyncSession = Depends(get_db),
    _=Depends(get_current_user),
):
    """Genera un Excel con el listado de ITV (RUA, fechas, estado, certificado, empresa)."""
    q = (
        select(ItvBus)
        .options(
            selectinload(ItvBus.bus).selectinload(Bus.asignaciones),
        )
    )
    filters = []
    if solo_vigentes:
        filters.append(ItvBus.es_vigente == True)
    if search:
        term = f"%{search.strip()}%"
        q = q.join(Bus, Bus.id_bus == ItvBus.id_bus)
        filters.append(
            or_(
                Bus.rua.ilike(term),
                ItvBus.numero_certificado.ilike(term),
            )
        )
    if filters:
        q = q.where(and_(*filters))

    q = q.order_by(ItvBus.fecha_vencimiento.asc())
    itv_list = (await db.execute(q)).scalars().unique().all()

    eot_ids = set()
    for itv in itv_list:
        bus = itv.bus
        if not bus:
            continue
        for a in bus.asignaciones:
            if a.fecha_fin_asignacion is None:
                eot_ids.add(a.id_eot)

    eot_nombres: dict[str, str] = {}
    if eot_ids:
        eot_rows = (
            await db.execute(
                select(Eot.id_eot_vmt_hex, Eot.eot_nombre).where(Eot.id_eot_vmt_hex.in_(eot_ids))
            )
        ).all()
        eot_nombres = {r[0]: r[1] for r in eot_rows}

    estado_filtro = estado.upper() if estado else None
    rows = []
    for itv in itv_list:
        est, dias = estado_itv(itv.fecha_vencimiento)
        if estado_filtro and est != estado_filtro:
            continue
        bus = itv.bus
        asig = next((a for a in (bus.asignaciones if bus else []) if a.fecha_fin_asignacion is None), None)
        empresa = "-"
        if asig:
            empresa = eot_nombres.get(asig.id_eot) or asig.id_eot or "-"
        rows.append({
            "rua": bus.rua if bus else "-",
            "numero_orden": bus.numero_orden if bus and bus.numero_orden is not None else "-",
            "empresa": empresa,
            "estado_bus": bus.estado_bus if bus else "-",
            "fecha_itv": itv.fecha_itv,
            "fecha_vencimiento": itv.fecha_vencimiento,
            "dias_para_vencer": dias if dias is not None else "-",
            "estado_itv": est,
            "resultado_itv": itv.resultado_itv or "-",
            "centro_itv": itv.centro_itv or "-",
            "numero_certificado": itv.numero_certificado or "-",
            "observaciones": itv.observaciones or "-",
            "es_vigente": "Sí" if itv.es_vigente else "No",
        })

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ITV"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    estado_fills = {
        "VIGENTE": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "POR_VENCER": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
        "CRITICO": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
        "VENCIDO": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    }

    headers = [
        ("rua", "RUA / Placa"),
        ("numero_orden", "N° Orden"),
        ("empresa", "Empresa"),
        ("estado_bus", "Estado Bus"),
        ("fecha_itv", "Fecha ITV"),
        ("fecha_vencimiento", "Fecha Vencimiento"),
        ("dias_para_vencer", "Días p/ vencer"),
        ("estado_itv", "Estado ITV"),
        ("resultado_itv", "Resultado"),
        ("centro_itv", "Centro ITV"),
        ("numero_certificado", "N° Certificado"),
        ("es_vigente", "Registro vigente"),
        ("observaciones", "Observaciones"),
    ]

    for col_num, (_, label) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_num, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = border_thin

    for row_idx, row_data in enumerate(rows, start=2):
        for col_num, (key, _) in enumerate(headers, start=1):
            val = row_data[key]
            if hasattr(val, "isoformat"):
                val = val.isoformat()
            cell = ws.cell(row=row_idx, column=col_num, value=val)
            cell.border = border_thin
            cell.alignment = align_center
            if key == "estado_itv":
                fill = estado_fills.get(str(val))
                if fill:
                    cell.fill = fill

    for col_idx in range(1, len(headers) + 1):
        max_len = 0
        for row_idx in range(1, (ws.max_row or 1) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    fecha = date.today().isoformat()
    return StreamingResponse(
        stream,
        headers={"Content-Disposition": f'attachment; filename="Reporte_ITV_{fecha}.xlsx"'},
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.get("/{id_itv}", response_model=ItvBusOut)
async def obtener_itv(id_itv: int, db: AsyncSession = Depends(get_db), _=Depends(get_current_user)):
    itv = (await db.execute(select(ItvBus).where(ItvBus.id_itv == id_itv))).scalar_one_or_none()
    if not itv:
        raise HTTPException(status_code=404, detail="Registro ITV no encontrado")
    est, dias = estado_itv(itv.fecha_vencimiento)
    return ItvBusOut(**{c.name: getattr(itv, c.name) for c in ItvBus.__table__.columns}, dias_para_vencer=dias, estado_itv=est)


@router.post("", response_model=ItvBusOut, status_code=status.HTTP_201_CREATED)
async def registrar_itv(
    body: ItvBusCreate,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_roles("ADMIN", "SUPERVISOR", "OPERADOR"))
):
    # Verificar que el bus existe
    bus = (await db.execute(select(Bus).where(Bus.id_bus == body.id_bus))).scalar_one_or_none()
    if not bus:
        raise HTTPException(status_code=404, detail="Bus no encontrado")

    # Si se registra una ITV marcada como vigente (default True), marcar las ITVs previas de ese bus como no vigentes
    if body.es_vigente:
        await db.execute(
            update(ItvBus)
            .where(ItvBus.id_bus == body.id_bus)
            .values(es_vigente=False)
        )

    itv = ItvBus(**body.model_dump())
    db.add(itv)
    await db.flush()
    await registrar_auditoria(
        db,
        accion="insert",
        tabla="itv_bus",
        usuario=user,
        registro_id=itv.id_itv or body.id_bus,
        datos_nuevos=body.model_dump(mode="json"),
        detalles=f"Alta ITV bus={body.id_bus}",
    )
    await db.commit()
    await db.refresh(itv)
    est, dias = estado_itv(itv.fecha_vencimiento)
    return ItvBusOut(**{c.name: getattr(itv, c.name) for c in ItvBus.__table__.columns}, dias_para_vencer=dias, estado_itv=est)


@router.put("/{id_itv}", response_model=ItvBusOut)
async def actualizar_itv(
    id_itv: int,
    body: ItvBusUpdate,
    db: AsyncSession = Depends(get_db),
    user=Depends(require_roles("ADMIN", "SUPERVISOR"))
):
    itv = (await db.execute(select(ItvBus).where(ItvBus.id_itv == id_itv))).scalar_one_or_none()
    if not itv:
        raise HTTPException(status_code=404, detail="Registro ITV no encontrado")

    update_data = body.model_dump(exclude_unset=True)

    # Si se actualiza es_vigente a True, asegurar que otras ITVs del bus pasen a False
    if update_data.get("es_vigente") is True:
        await db.execute(
            update(ItvBus)
            .where(and_(ItvBus.id_bus == itv.id_bus, ItvBus.id_itv != id_itv))
            .values(es_vigente=False)
        )

    for k, v in update_data.items():
        setattr(itv, k, v)
    await registrar_auditoria(
        db,
        accion="update",
        tabla="itv_bus",
        usuario=user,
        registro_id=id_itv,
        datos_nuevos=update_data,
        detalles=f"Actualización ITV id={id_itv}",
    )
    await db.commit()
    await db.refresh(itv)
    est, dias = estado_itv(itv.fecha_vencimiento)
    return ItvBusOut(**{c.name: getattr(itv, c.name) for c in ItvBus.__table__.columns}, dias_para_vencer=dias, estado_itv=est)


@router.get("/historial/{id_bus}", response_model=list[HistorialItvOut])
async def historial_itv_bus(id_bus: int, db: AsyncSession = Depends(get_db), _=Depends(get_current_user)):
    # Orden cronológico por vencimiento, luego registro (para deduplicar fechas iguales)
    records = (await db.execute(
        select(ItvBus)
        .where(ItvBus.id_bus == id_bus)
        .order_by(ItvBus.fecha_vencimiento.asc(), ItvBus.fecha_registro.asc(), ItvBus.id_itv.asc())
    )).scalars().all()

    # Una fila por (fecha_itv, fecha_vencimiento): conservar la más antigua
    seen_keys: set[tuple] = set()
    unique: list[ItvBus] = []
    for rec in records:
        key = (rec.fecha_itv, rec.fecha_vencimiento)
        if key in seen_keys:
            continue
        seen_keys.add(key)
        unique.append(rec)

    historial = []
    for idx, rec in enumerate(unique):
        prev_venc = unique[idx - 1].fecha_vencimiento if idx > 0 else None
        diff_days = (rec.fecha_itv - prev_venc).days if prev_venc and rec.fecha_itv else None

        historial.append(HistorialItvOut(
            id_historial=rec.id_itv,
            id_bus=rec.id_bus,
            fecha_vencimiento_anterior=prev_venc,
            fecha_itv_actual=rec.fecha_itv,
            fecha_vencimiento_actual=rec.fecha_vencimiento,
            diferencia_dias=diff_days,
            observaciones=rec.observaciones,
            fecha_registro=rec.fecha_registro
        ))

    return list(reversed(historial))
